"""
Handles document uploads and text extraction for AI analysis.
"""
import os
import uuid
from datetime import datetime
from typing import Optional, BinaryIO
from pathlib import Path

from app.database import get_db
from app.models.workload_intelligence import WorkloadDocument, EvaluationDocument


UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {".docx", ".pdf", ".xlsx", ".txt", ".csv"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB


class DocumentService:
    """
    Handles document upload, storage, and text extraction.
    """
    
    async def upload_workload_document(
        self, 
        workload_id: int, 
        file: BinaryIO,
        filename: str,
        uploaded_by: str
    ) -> WorkloadDocument:
        """Upload a document for a workload."""
        
        # Validate
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError(f"File type {ext} not allowed")
        
        # Generate unique filename
        unique_name = f"{uuid.uuid4()}{ext}"
        file_path = UPLOAD_DIR / "workloads" / str(workload_id) / unique_name
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        # Save file
        content = file.read()
        if len(content) > MAX_FILE_SIZE:
            raise ValueError(f"File too large (max {MAX_FILE_SIZE // 1024 // 1024}MB)")
        
        with open(file_path, 'wb') as f:
            f.write(content)
        
        # Extract text
        extracted_text = self._extract_text(file_path, ext)
        
        # Create record
        with get_db() as db:
            doc = WorkloadDocument(
                workload_id=workload_id,
                filename=unique_name,
                original_filename=filename,
                file_path=str(file_path),
                file_type=ext[1:],  # Remove dot
                file_size_bytes=len(content),
                extracted_text=extracted_text,
                extraction_status="completed" if extracted_text else "failed",
                uploaded_by=uploaded_by
            )
            db.add(doc)
            db.commit()
            db.refresh(doc)
            return doc
    
    async def upload_evaluation_document(
        self,
        evaluation_id: int,
        file: BinaryIO,
        filename: str,
        document_type: str,
        uploaded_by: str
    ) -> EvaluationDocument:
        """Upload a document for an evaluation."""
        
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError(f"File type {ext} not allowed")
        
        unique_name = f"{uuid.uuid4()}{ext}"
        file_path = UPLOAD_DIR / "evaluations" / str(evaluation_id) / unique_name
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        content = file.read()
        if len(content) > MAX_FILE_SIZE:
            raise ValueError(f"File too large")
        
        with open(file_path, 'wb') as f:
            f.write(content)
        
        extracted_text = self._extract_text(file_path, ext)
        
        with get_db() as db:
            doc = EvaluationDocument(
                evaluation_id=evaluation_id,
                filename=unique_name,
                original_filename=filename,
                file_path=str(file_path),
                file_type=ext[1:],
                file_size_bytes=len(content),
                document_type=document_type,
                extracted_text=extracted_text,
                uploaded_by=uploaded_by
            )
            db.add(doc)
            db.commit()
            db.refresh(doc)
            return doc
    
    def _extract_text(self, file_path: Path, ext: str) -> Optional[str]:
        """Extract text from document for AI analysis."""
        try:
            if ext == ".docx":
                return self._extract_docx(file_path)
            elif ext == ".pdf":
                return self._extract_pdf(file_path)
            elif ext == ".xlsx":
                return self._extract_xlsx(file_path)
            elif ext in [".txt", ".csv"]:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            return None
        except Exception as e:
            print(f"Text extraction failed: {e}")
            return None
    
    def _extract_docx(self, file_path: Path) -> str:
        """Extract text from Word document."""
        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except ImportError:
            return "[python-docx not installed - cannot extract Word documents]"
        except Exception as e:
            return f"[Error extracting Word document: {e}]"
    
    def _extract_pdf(self, file_path: Path) -> str:
        """Extract text from PDF."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            text = []
            for page in doc:
                text.append(page.get_text())
            return "\n".join(text)
        except ImportError:
            return "[PyMuPDF not installed - cannot extract PDF documents]"
        except Exception as e:
            return f"[Error extracting PDF: {e}]"
    
    def _extract_xlsx(self, file_path: Path) -> str:
        """Extract text from Excel (as CSV-like format)."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = []
            for sheet in wb.worksheets:
                text.append(f"=== Sheet: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) if cell is not None else "" for cell in row]
                    text.append(",".join(row_text))
            return "\n".join(text)
        except ImportError:
            return "[openpyxl not installed - cannot extract Excel documents]"
        except Exception as e:
            return f"[Error extracting Excel: {e}]"


# Singleton instance
document_service = DocumentService()
